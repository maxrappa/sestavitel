import openpyxl
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from datetime import datetime
from tkinter import Tk, filedialog
import os

Tk().withdraw()

input_file = filedialog.askopenfilename(
    title="Select Excel file",
    filetypes=[("Excel files", "*.xlsx")]
)

if not input_file:
    print("No file selected.")
    exit()

wb = load_workbook(input_file)
ws = wb.active

# CLEANING
if 'A3:C3' in ws.merged_cells: 
    ws.unmerge_cells('A3:C3')

if 'A1:C1' in ws.merged_cells:
    ws.unmerge_cells('A1:C1')

ws.delete_rows(2, 2)
ws.delete_rows(ws.max_row)
ws['B1'] = 'Nazev'
ws['C1'] = 'Skupina'

# RERANGE
header_row = 1
current_columns = {}
for cell in ws[header_row]:
    if cell.value is not None:
        clean_name = ' '.join(str(cell.value).split()).strip()
        if clean_name in current_columns:
            current_columns[clean_name].append(cell.column)
        else:
            current_columns[clean_name] = [cell.column]

desired_order = [
    'Sortiment',
    'Foto 1',
    'Foto 2',
    'Nazev',
    'Skupina',
    'Prodejní cena 2024-Y24',
    'Prodejní cena 2025-Y25',
    'Prodejní cena 2026-Y26',
    'Nákupní cena',
    'Skut. cena / kus',
    'VOZ CENÍK',
    'Marže 2024-Y24',
    'Marže 2025-Y25',
    'Marže 2026-Y26',
    'Marže % 2024-Y24',
    'Marže % 2025-Y25',
    'Marže % 2026-Y26',
    "Množství 2024-Y24",
    "Množství 2025-Y25",
    "Množství 2026-Y26",
    'Příjem obd. 2024-Y24',
    'Příjem obd. 2025-Y25',
    'Příjem obd. 2026-Y26',
    'Sklade fyzicky',
    'Skladem volné',
    'Suma zbývá dodat',
    '(Skladem + zbývá dodat) / prodej 2025',
    'Order ctn',
    'Order pcs',
    'Komentař Jarmila',
    'Komentař nakup',
    'Balení 3',
    'Minimální odběr pro OV',
    'Objem',
    'Objem celkem',
    'Poznámka',
    'Prim. dodavatel',
    'EAN'
]

data_rows = list(ws.iter_rows(min_row=2, max_row=ws.max_row))

# SORTING
sort_col_name = 'Prodejní cena 2025-Y25'
old_sort_col_idx = current_columns.get(sort_col_name, [])

if old_sort_col_idx:
    old_sort_col_idx = old_sort_col_idx[0]
    def get_sort_value(row):
        val = row[old_sort_col_idx - 1].value
        try: 
            return float(val) if val is not None else 0
        except ValueError:
            return 0
    sorted_data_rows = sorted(data_rows, key=get_sort_value, reverse=True)
else:
    sorted_data_rows = data_rows

new_ws = wb.create_sheet(title='Reranged')

new_header_row = 2
for new_col_idx, col_name in enumerate(desired_order, start=1):
    new_ws.cell(row=new_header_row, column=new_col_idx, value=col_name)

# MAPPING FOR CALCULATIONS
def find_col_letter(name):
    return get_column_letter(desired_order.index(name) + 1)

l_sklad_vol = find_col_letter('Skladem volné')
l_sum_zbyva = find_col_letter('Suma zbývá dodat')
l_mnoz_2025 = find_col_letter('Množství 2025-Y25')
l_order_ctn = find_col_letter('Order ctn')
l_baleni_3  = find_col_letter('Balení 3')
l_objem     = find_col_letter('Objem')

for new_row_idx, row_cells in enumerate(sorted_data_rows, start=3):
    for new_col_idx, col_name in enumerate(desired_order, start=1):
        old_value = None

        if col_name == 'Order ctn': 
            new_ws.cell(row=new_row_idx, column=new_col_idx, value=0)
        if col_name == 'Suma zbývá dodat':
            raw_del_indices = current_columns.get('Zbývá dodat ks', [])
            if raw_del_indices:
                plus_components = [f"'{ws.title}'!{get_column_letter(idx)}{row_cells[0].row}" for idx in raw_del_indices]
                new_ws.cell(row=new_row_idx, column=new_col_idx, value=f"={'+'.join(plus_components)}")
            else:
                new_ws.cell(row=new_row_idx, column=new_col_idx, value=0)

        else:
            old_indices = current_columns.get(col_name, [])
            if old_indices:
                old_value = row_cells[old_indices[0] - 1].value

                if col_name in ['Marže % 2024-Y24', 'Marže % 2025-Y25', 'Marže % 2026-Y26'] and isinstance(old_value, (int, float)):
                    old_value = old_value / 100

            new_ws.cell(row=new_row_idx, column=new_col_idx, value=old_value)

    new_ws.cell(row=new_row_idx, column=desired_order.index('(Skladem + zbývá dodat) / prodej 2025') + 1, 
                    value=f'=IFERROR(({l_sklad_vol}{new_row_idx}+{l_sum_zbyva}{new_row_idx})/{l_mnoz_2025}{new_row_idx}, 0)')
        
    new_ws.cell(row=new_row_idx, column=desired_order.index('Order pcs') + 1, 
                value=f'={l_order_ctn}{new_row_idx}*{l_baleni_3}{new_row_idx}')
    
    new_ws.cell(row=new_row_idx, column=desired_order.index('Objem celkem') + 1, 
                value=f'={l_objem}{new_row_idx}*{l_order_ctn}{new_row_idx}')

raw_delivery_indices = current_columns.get('Zbývá dodat ks', [])
raw_date_indices = current_columns.get('Datum dodání z.', [])
append_start_col = len(desired_order) + 1

if raw_delivery_indices:
        for d_idx, raw_qty_col in enumerate(raw_delivery_indices):

            qty_col_idx = append_start_col + (d_idx * 2)
            date_col_idx = qty_col_idx + 1

            if d_idx < len(raw_date_indices):
                raw_date_col = raw_date_indices[d_idx]
            else:
                raw_date_col = raw_qty_col + 1

            new_ws.cell(row=new_header_row, column=qty_col_idx, value=f"Zbyva dodat {d_idx+1}")
            new_ws.cell(row=new_header_row, column=date_col_idx, value=f"Datum dodáni {d_idx+1}")

            for r_idx, row_cells in enumerate(sorted_data_rows, start=3):
                qty_val = row_cells[raw_qty_col - 1].value

                if qty_val == 0 or qty_val is None:
                    new_ws.cell(row=r_idx, column=qty_col_idx, value=None)
                    new_ws.cell(row=r_idx, column=date_col_idx, value=None)
                    continue
                
                new_ws.cell(row=r_idx, column=qty_col_idx, value=qty_val)

                date_val = row_cells[raw_date_col - 1].value
                new_ws.cell(row=r_idx, column=date_col_idx, value=date_val)

objem_celkem_col_idx = desired_order.index('Objem celkem') + 1
objem_letter = get_column_letter(objem_celkem_col_idx)

last_data_row = len(sorted_data_rows) + 2
total_cbm_formula = f'=SUM({objem_letter}3:{objem_letter}{last_data_row})'

new_ws.cell(row=1, column=28, value='Total CBM')
new_ws.cell(row=1, column=29, value=total_cbm_formula)

# STYLING

COLORS = {
    'red': PatternFill(start_color='d4ebff', end_color='d4ebff', fill_type='solid'),
    'yellow': PatternFill(start_color='fbff1f', end_color='fbff1f', fill_type='solid'),
    'blue': PatternFill(start_color='1c86ff', end_color='1c86ff', fill_type='solid'),
    'violet': PatternFill(start_color='bb1cff', end_color='bb1cff', fill_type='solid'),
    'light_blue': PatternFill(start_color='c9e6ff', end_color='c9e6ff', fill_type='solid'),
    'light_green': PatternFill(start_color='d0ffb3', end_color='d0ffb3', fill_type='solid'),
    'light_yellow': PatternFill(start_color='ffe496', end_color='ffe496', fill_type='solid'),
    'soft_gray': PatternFill(start_color='d1d1d1', end_color='d1d1d1', fill_type='solid'),
    'soft_orange': PatternFill(start_color='ffd5bd', end_color='ffd5bd', fill_type='solid'),
}

base_font = Font(name='Segoe UI', size=12, bold=False)
header_font = Font(name='Segoe UI', size=9, bold=True)
important_font = Font(name='Segoe UI', size=14, bold=True)

new_ws.row_dimensions[1].height = 25
new_ws.row_dimensions[2].height = 40

new_ws.column_dimensions['I'].hidden = True
new_ws.column_dimensions['J'].hidden = True

for row_idx in range(3, new_ws.max_row + 1):
    new_ws.row_dimensions[row_idx].height = 18

for col in new_ws.columns:
    col_letter = get_column_letter(col[0].column)

    max_len = 0
    for cell in col[2:]:
        if cell.value is not None:
            max_len = max(max_len, len(str(cell.value)))
            
    new_ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

    calculated_width = max_len + 3

    if calculated_width > 35:
        final_width = 35
    elif calculated_width < 11:
        final_width = 11
    else:
        final_width = calculated_width

    new_ws.column_dimensions[col_letter].width = final_width

# COLORING
for row_idx in range(1, new_ws.max_row + 1): 
    for col_idx in range(1, new_ws.max_column + 1):
        cell = new_ws.cell(row=row_idx, column=col_idx)

        cell.font = base_font

        if 6 <= col_idx <= 8:
            cell.fill = COLORS['light_blue']

        if 12 <= col_idx <=17:
            cell.fill = COLORS['soft_orange']
        
        if 18 <= col_idx <= 20:
            cell.fill = COLORS['light_yellow']

        if 21 <= col_idx <= 23:
            cell.fill = COLORS['soft_gray']

        if 28 <= col_idx <= 29:
            cell.fill = COLORS['yellow']
            cell.font = Font(name='Segoe UI', size=14, bold=True, color='ff3838')
        
        if col_idx == 30:
            cell.fill = COLORS['violet']
            cell.font = Font(name='Segoe UI', size=10, bold=True, color='fbff1f')
        
        if col_idx == 31:
            cell.fill = COLORS['blue']
            cell.font = Font(name='Segoe UI', size=10, bold=True, color='fbff1f')

# CUSTOMS
for row_idx in range(1, new_ws.max_row + 1): 
    for col_idx in range(1, new_ws.max_column + 1):
        cell = new_ws.cell(row=row_idx, column=col_idx)
        
        cell.number_format = '###,##0.00'

        thin_side = Side(border_style='thin', color='242424')
        cell.border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        new_ws.column_dimensions['AL'].width = 25

        if row_idx == 1 and col_idx == 1:
            new_ws.merge_cells('A1:D1')
            cell.value = 'Nazev sestavu'
            cell.font = Font(name='Segoe UI', size=16, bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')

        if row_idx == 2:
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        
        if row_idx == 1:
            cell.fill = COLORS['light_green']
            cell.font = header_font
            cell.number_format = '0.00'

        if row_idx == 1 and col_idx == 7: 
            cell.value = 'RAZENO!'
            cell.fill = COLORS['yellow']
            cell.font = Font(name='Segoe UI', size=12, bold=True, color='ff3838')
            cell.alignment = Alignment(horizontal='center', vertical='center')

        if 16 <= col_idx <= 33 or col_idx == 10 or col_idx == 11:
            cell.number_format = '#,##0'
            
        if 15 <= col_idx <=17 or col_idx == 27:
            cell.number_format = '0.0%'
        
        if row_idx == 1 and col_idx == 8: 
            cell.value = datetime.now().date()
            cell.number_format = 'dd.mm.yyyy'
            cell.font = Font(name='Segoe UI', size=10, bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')

        if col_idx == 39 or col_idx == 41:
            cell.number_format = '#,##0'

sales_range = new_ws['E1:F1']
for row in sales_range:   
    for cell in row:

        highlight_fill = PatternFill(start_color='a6d5ff', end_color='a6d5ff', fill_type='solid')
        highlight_font = Font(name='Segoe UI', size=12, bold=True)

        cell.fill = highlight_fill
        cell.font = highlight_font

        if cell.column == 5:
            cell.value = 'Vic než:'

        if cell.column == 6:        
            dynamic_rule = CellIsRule(
                operator='greaterThan',
                formula=['=$F$1'],
                stopIfTrue=True,
                fill=highlight_fill,
                font=highlight_font
            )

            sales_data_range = f'F3:H{new_ws.max_row}'
            new_ws.conditional_formatting.add(sales_data_range, dynamic_rule)

            cell.value = 'napiš tady'
            cell.number_format = '#,##0'       
    
# TOTAL CBM
total_cbm_range = new_ws['AB1:AC1']
for row in total_cbm_range:
    for cell in row: 
        if cell.column == 29:
            cell.fill = COLORS['blue']
            cell.font = Font(name='Segoe UI', size=16, bold=True, color='fbff1f')
            cell.alignment = Alignment(horizontal='right', vertical='center', wrap_text=True)
            cell.number_format = '0.00'

        else:
            cell.fill = COLORS['blue']
            cell.font = Font(name='Segoe UI', size=11, bold=True, color='fbff1f')
            cell.alignment = Alignment(horizontal='right', vertical='center')

folder = os.path.dirname(input_file)
filename = os.path.splitext(os.path.basename(input_file))[0]

output_file = os.path.join(folder, f"{filename}_processed.xlsx")

wb.save(output_file)

print(f"Saved to:\n{output_file}")

print('worked out')